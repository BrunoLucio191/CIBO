#!/usr/bin/env python3
"""Calendário de conteúdo -> planilha de produção (uma linha por conteúdo).

Uso: python3 calendario_xlsx.py calendario.json calendario.xlsx
calendario.json: lista de objetos com as chaves de COLS (as de produção podem faltar).
Requer openpyxl (pip install -r requirements.txt).
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLS = [('data', 'Data', 12), ('plataforma', 'Plataforma', 12), ('formato', 'Formato', 14),
        ('pilar', 'Pilar', 18), ('funil', 'Etapa do funil', 12), ('tema', 'Tema', 32),
        ('gancho', 'Gancho (1ª frase / 3s)', 40), ('estrutura', 'Estrutura do roteiro / slides', 60),
        ('cta', 'CTA', 28), ('metrica', 'Métrica de sucesso', 26), ('origem', 'Reaproveitado de', 16),
        ('status', 'Status', 12), ('responsavel', 'Responsável', 14), ('link', 'Link', 24)]

rows = json.load(open(sys.argv[1], encoding='utf-8'))
wb = Workbook(); ws = wb.active; ws.title = 'Calendário'
ws.append([c[1] for c in COLS])
for r in rows:
    ws.append(['\n'.join(r[k]) if isinstance(r.get(k), list) else r.get(k, '') for k, _, _ in COLS])
head = PatternFill('solid', fgColor='1F2A44')
for i, (_, _, w) in enumerate(COLS, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
    c = ws.cell(1, i); c.font = Font(bold=True, color='FFFFFF'); c.fill = head
    c.alignment = Alignment(vertical='center', wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(vertical='top', wrap_text=True)
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
wb.save(sys.argv[2]); print(f'{len(rows)} conteúdos -> {sys.argv[2]}')
